"""Atomic, non-destructive write-back to the source workbook."""
from __future__ import annotations

import os
import shutil
import sqlite3
import tempfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .reader import KNOWN_HEADERS, norm_header, parse_date_cell


class SheetLockedError(Exception):
    pass


@dataclass
class Update:
    sheet_row: int
    last_advertised: date | None = None
    replied: bool | None = None
    status: str | None = None
    last_error: str | None = None
    attempts: int | None = None


@dataclass
class Change:
    row: int
    column: str
    old: Any
    new: Any

    def describe(self) -> str:
        return f"row {self.row}: {self.column} {self.old!r} -> {self.new!r}"


@dataclass
class WriteResult:
    path: Path
    changes: list[Change]
    rows_touched: int

    @property
    def summary(self) -> str:
        return f"{len(self.changes)} cell(s) updated across {self.rows_touched} row(s)"


def lockfile_for(path: Path) -> Path:
    return path.with_name("~$" + path.name)


def assert_not_locked(path: Path) -> None:
    lock = lockfile_for(path)
    if lock.exists():
        raise SheetLockedError(
            f"{path.name} is open in Excel (lock file {lock.name} present). "
            "Close it before writing results."
        )


def _header_map(ws) -> tuple[int, dict[str, int], list[str]]:
    header_row = None
    headers: list[str] = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20),
                                             values_only=True), start=1):
        if any(v is not None and str(v).strip() for v in row):
            header_row = r_idx
            headers = ["" if v is None else str(v).strip() for v in row]
            break
    if header_row is None:
        return 1, {}, []
    cols: dict[str, int] = {}
    for idx, name in enumerate(headers, start=1):
        key = KNOWN_HEADERS.get(norm_header(name))
        if key and key not in cols:
            cols[key] = idx
    return header_row, cols, headers


def collect_updates(conn: sqlite3.Connection, run_id: int | None = None,
                    only_real_sends: bool = True) -> list[Update]:
    """Build write-back rows from SQLite (never from a live sheet read)."""
    where = ["1=1"]
    params: list[Any] = []
    if only_real_sends:
        where.append("s.dry_run=0")
    if run_id is not None:
        where.append("s.run_id=?")
        params.append(run_id)
    rows = conn.execute(
        f"""SELECT c.id AS contact_id, c.sheet_row, c.last_advertised, c.replied,
                   c.status, c.last_error, c.attempts,
                   SUM(CASE WHEN s.state='sent' THEN 1 ELSE 0 END) AS sent_count,
                   MAX(s.sent_at) AS last_sent
            FROM sends s JOIN contacts c ON c.id=s.contact_id
            WHERE {' AND '.join(where)}
            GROUP BY c.id ORDER BY c.sheet_row""",
        params,
    )
    out: list[Update] = []
    for r in rows:
        if not r["sheet_row"]:
            continue
        last = r["last_advertised"]
        sent = int(r["sent_count"] or 0) > 0
        out.append(Update(
            sheet_row=int(r["sheet_row"]),
            # LastAdvertised is only nudged for rows that actually sent.
            last_advertised=(date.fromisoformat(str(last)[:10])
                             if sent and last else None),
            replied=None,
            status=r["status"] or ("sent" if sent else None),
            last_error=r["last_error"],
            attempts=r["attempts"],
        ))
    return out


def _existing_date(cell, date_order: str = "DMY") -> date | None:
    value = cell.value
    if value is None:
        return None
    parsed, _, _ = parse_date_cell(value, date_order)
    return parsed


def plan(path: Path, updates: Iterable[Update], tab: str | None = None,
         date_order: str = "DMY") -> list[Change]:
    """Dry-run of the write: returns the exact cells that would change."""
    p = Path(path)
    wb = load_workbook(p, data_only=False, keep_vba=p.suffix.lower() == ".xlsm")
    try:
        ws = wb[tab] if tab else wb.active
        _, cols, _ = _header_map(ws)
        changes: list[Change] = []
        for u in updates:
            r = u.sheet_row
            if not r or r < 2 or r > ws.max_row:
                continue
            if "last_advertised" in cols and u.last_advertised:
                cell = ws.cell(row=r, column=cols["last_advertised"])
                old = _existing_date(cell, date_order)
                if old is None or u.last_advertised > old:
                    changes.append(Change(r, "LastAdvertised",
                                          old.isoformat() if old else None,
                                          u.last_advertised.isoformat()))
            if "replied" in cols and u.replied:
                cell = ws.cell(row=r, column=cols["replied"])
                if not (cell.value is True or str(cell.value).strip().lower() in
                        ("1", "y", "yes", "true", "x")):
                    changes.append(Change(r, "Replied", cell.value, "TRUE"))
            if "status" in cols and u.status:
                cell = ws.cell(row=r, column=cols["status"])
                if cell.value != u.status:
                    changes.append(Change(r, "Status", cell.value, u.status))
            if "last_error" in cols and u.last_error is not None:
                cell = ws.cell(row=r, column=cols["last_error"])
                if cell.value != u.last_error:
                    changes.append(Change(r, "LastError", cell.value, u.last_error))
            if "attempts" in cols and u.attempts is not None:
                cell = ws.cell(row=r, column=cols["attempts"])
                if cell.value != u.attempts:
                    changes.append(Change(r, "Attempts", cell.value, u.attempts))
        return changes
    finally:
        wb.close()


def write_back(path: str | Path, updates: Iterable[Update], tab: str | None = None,
               backup: bool = True, date_order: str = "DMY") -> WriteResult:
    """Write to a temp file in the same directory, then os.replace() over the original."""
    p = Path(path)
    updates = [u for u in updates if u.sheet_row]
    for u in updates:
        if not u.sheet_row or u.sheet_row < 2:
            raise ValueError(f"refusing to write to row {u.sheet_row}")
    assert_not_locked(p)

    # keep_vba keeps a macro-enabled workbook macro-enabled instead of silently
    # stripping vbaProject.bin on save.
    wb = load_workbook(p, data_only=False, keep_vba=p.suffix.lower() == ".xlsm")
    tmp_name = None
    try:
        ws = wb[tab] if tab else wb.active
        _, cols, _ = _header_map(ws)
        changes: list[Change] = []
        touched: set[int] = set()

        for u in updates:
            r = u.sheet_row
            if r > ws.max_row:
                continue
            if "last_advertised" in cols and u.last_advertised:
                cell = ws.cell(row=r, column=cols["last_advertised"])
                old = _existing_date(cell, date_order)
                if old is None or u.last_advertised > old:
                    cell.value = u.last_advertised
                    cell.number_format = "yyyy-mm-dd"
                    changes.append(Change(r, "LastAdvertised",
                                          old.isoformat() if old else None,
                                          u.last_advertised.isoformat()))
                    touched.add(r)
            if "replied" in cols and u.replied:
                cell = ws.cell(row=r, column=cols["replied"])
                current = str(cell.value).strip().lower()
                if cell.value is not True and current not in ("1", "y", "yes", "true", "x"):
                    changes.append(Change(r, "Replied", cell.value, True))
                    cell.value = True
                    touched.add(r)
            if "status" in cols and u.status and ws.cell(row=r, column=cols["status"]).value != u.status:
                cell = ws.cell(row=r, column=cols["status"])
                changes.append(Change(r, "Status", cell.value, u.status))
                cell.value = u.status
                touched.add(r)
            if "last_error" in cols and u.last_error is not None:
                cell = ws.cell(row=r, column=cols["last_error"])
                if cell.value != u.last_error:
                    changes.append(Change(r, "LastError", cell.value, u.last_error))
                    cell.value = u.last_error
                    touched.add(r)
            if "attempts" in cols and u.attempts is not None:
                cell = ws.cell(row=r, column=cols["attempts"])
                if cell.value != u.attempts:
                    changes.append(Change(r, "Attempts", cell.value, u.attempts))
                    cell.value = u.attempts
                    touched.add(r)

        if backup and p.exists():
            try:
                shutil.copy2(p, p.with_name(p.stem + ".dispatch-backup" + p.suffix))
            except OSError:
                pass

        fd, tmp_name = tempfile.mkstemp(prefix=p.stem + ".", suffix=p.suffix, dir=str(p.parent))
        os.close(fd)
        wb.save(tmp_name)
        wb.close()
        os.replace(tmp_name, p)
        tmp_name = None
        return WriteResult(p, changes, len(touched))
    finally:
        try:
            wb.close()
        except Exception:
            pass
        if tmp_name and os.path.exists(tmp_name):
            try:
                os.unlink(tmp_name)
            except OSError:
                pass


def flush_run(conn: sqlite3.Connection, cfg, run_id: int | None) -> WriteResult:
    updates = collect_updates(conn, run_id)
    return write_back(cfg.sheet_file, updates, cfg.campaign.sheet_tab,
                      date_order=cfg.campaign.date_order)
