"""Spreadsheet round-trip and write-back safety (section 15.7)."""
from __future__ import annotations

from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from dispatch import db
from dispatch.engine import preflight
from dispatch.sheet import writer as sheet_writer
from dispatch.sheet.reader import read_sheet
from tests.conftest import HEADERS


def test_flush_updates_only_sent_rows_and_preserves_everything_else(campaign):
    headers = HEADERS + ["CustomField"]
    contacts = [
        {"email": "sent@example.com", "first_name": "Sent", "CustomField": "keep-me"},
        {"email": "failed@example.com", "first_name": "Failed", "CustomField": "also-keep"},
    ]
    cfg = campaign(contacts, headers=headers)
    conn = db.connect(cfg.db_path)
    db.init_db(conn)
    preflight.import_sheet(cfg, conn)

    ids = [int(r["id"]) for r in db.contacts_in_sheet(conn)]
    run_id = db.create_run(conn, 2, False)
    db.enqueue(conn, run_id, ids, "sam")

    first = db.claim_next(conn, run_id)
    db.mark_sent(conn, first["send_id"], "250", dry=False)
    second = db.claim_next(conn, run_id)
    db.mark_failed(conn, second["send_id"], "failed_hard", "550", "no such user")

    result = sheet_writer.flush_run(conn, cfg, run_id)
    conn.close()

    assert result.changes
    wb = load_workbook(cfg.sheet_file)
    ws = wb["Leads"]
    cols = {ws.cell(row=1, column=i).value: i for i in range(1, ws.max_column + 1)}

    sent_row, failed_row = 2, 3
    assert ws.cell(row=sent_row, column=cols["LastAdvertised"]).value is not None
    assert str(ws.cell(row=sent_row, column=cols["LastAdvertised"]).value)[:10] == \
        date.today().isoformat()
    assert ws.cell(row=failed_row, column=cols["LastAdvertised"]).value is None
    assert ws.cell(row=failed_row, column=cols["Status"]).value == "failed_hard"
    assert ws.cell(row=failed_row, column=cols["LastError"]).value == "no such user"

    # unlisted column + formatting survive
    assert ws.cell(row=sent_row, column=cols["CustomField"]).value == "keep-me"
    assert ws.cell(row=failed_row, column=cols["CustomField"]).value == "also-keep"
    assert ws.cell(row=1, column=1).font.bold is True
    assert ws.cell(row=2, column=1).fill.start_color.rgb.endswith("FFF2CC")
    wb.close()


def test_last_advertised_never_moves_backwards(campaign):
    future = date.today() + timedelta(days=60)
    contacts = [{"email": "future@example.com", "last_advertised": future}]
    cfg = campaign(contacts)
    conn = db.connect(cfg.db_path)
    db.init_db(conn)
    preflight.import_sheet(cfg, conn)

    ids = [int(r["id"]) for r in db.contacts_in_sheet(conn)]
    run_id = db.create_run(conn, 1, False)
    db.enqueue(conn, run_id, ids, "sam")
    claimed = db.claim_next(conn, run_id)
    db.mark_sent(conn, claimed["send_id"], "250", dry=False)
    row = conn.execute("SELECT last_advertised FROM contacts WHERE id=?", (ids[0],)).fetchone()
    assert row["last_advertised"] == future.isoformat()

    sheet_writer.flush_run(conn, cfg, run_id)
    conn.close()

    data = read_sheet(cfg.sheet_file)
    assert data.contacts[0].last_advertised == future


def test_refuses_to_write_when_excel_lockfile_exists(campaign):
    cfg = campaign([{"email": "a@example.com"}])
    lock = cfg.sheet_file.with_name("~$" + cfg.sheet_file.name)
    lock.write_text("locked")
    try:
        with pytest.raises(sheet_writer.SheetLockedError):
            sheet_writer.write_back(cfg.sheet_file, [])
    finally:
        lock.unlink()


def test_planned_changes_are_a_dry_run(campaign):
    contacts = [{"email": "a@example.com"}]
    cfg = campaign(contacts)
    conn = db.connect(cfg.db_path)
    db.init_db(conn)
    preflight.import_sheet(cfg, conn)
    ids = [int(r["id"]) for r in db.contacts_in_sheet(conn)]
    run_id = db.create_run(conn, 1, False)
    db.enqueue(conn, run_id, ids, "sam")
    claimed = db.claim_next(conn, run_id)
    db.mark_sent(conn, claimed["send_id"], "250", dry=False)

    before = cfg.sheet_file.read_bytes()
    changes = sheet_writer.plan(cfg.sheet_file, sheet_writer.collect_updates(conn, run_id))
    assert cfg.sheet_file.read_bytes() == before
    assert any(c.column == "LastAdvertised" for c in changes)
    conn.close()
