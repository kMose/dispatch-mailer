"""Adversarial tests written during the bug audit.

Each of these started life as a failing test; they guard the fixes.
"""
from __future__ import annotations

import time
from datetime import date, datetime
from zoneinfo import ZoneInfo

import pytest

from dispatch import db
from dispatch.config import FullConfig
from dispatch.engine import feasibility, preflight
from dispatch.engine.worker import Worker
from dispatch.mail import blocks, monitor
from dispatch.sheet import writer as sheet_writer


def prepare(cfg):
    conn = db.connect(cfg.db_path)
    db.init_db(conn)
    preflight.import_sheet(cfg, conn)
    conn.close()


# -- 1. pause must not squeeze one more message out ------------------------
def test_pause_during_countdown_sends_nothing(campaign):
    contacts = [{"email": f"pause{i}@example.com"} for i in range(3)]
    cfg = campaign(contacts, quota=3, dry_run=True, fixed_seconds=3)
    prepare(cfg)
    events: list[dict] = []
    worker = Worker(cfg, quota=3, dry_run=True, on_event=events.append, prepared=True)
    worker.start()
    try:
        deadline = time.time() + 15
        while time.time() < deadline and worker.stats.sent < 1:
            time.sleep(0.05)
        worker.pause()
        snapshot = worker.stats.sent
        time.sleep(4.0)  # longer than the 3s interval
        assert worker.stats.sent == snapshot, (
            f"pause let another send through: {snapshot} -> {worker.stats.sent}"
        )
        worker.resume()
    finally:
        worker.stop()
        worker.join(15)
    assert not worker.is_alive()


# -- 2. a window that opens >24h away is a wait, not a closed window ------
def test_window_further_than_a_day_is_waited_for(campaign):
    contacts = [{"email": "window@example.com"}]
    cfg = campaign(contacts, quota=1, dry_run=True, window_enabled=True)
    window = cfg.campaign.send_window
    names = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    now = datetime.now(window.tz)
    # pick a day two days out so the next opening is >24h away
    window.days = [names[(now.weekday() + 2) % 7]]
    window.start = "09:00"
    window.end = "10:00"
    prepare(cfg)

    events: list[dict] = []
    worker = Worker(cfg, quota=1, dry_run=True, on_event=events.append, prepared=True)
    worker.start()
    try:
        deadline = time.time() + 10
        while time.time() < deadline and not any(e.get("type") == "window" for e in events):
            time.sleep(0.05)
        assert any(e.get("type") == "window" for e in events), "no window event"
        time.sleep(0.5)
        assert worker.is_alive(), f"worker stopped instead of waiting: {worker.stop_reason}"
    finally:
        worker.stop()
        worker.join(15)


# -- 3. a rate-limited row is failed_soft, not failed_hard ----------------
def test_rate_limit_event_reports_failed_soft(campaign, smtp_sink):
    smtp_sink.fail_next = 1
    cfg = campaign([{"email": "rl@example.com"}], quota=1, smtp_port=smtp_sink.port)
    prepare(cfg)
    events: list[dict] = []
    worker = Worker(cfg, quota=1, dry_run=False, on_event=events.append, prepared=True)
    worker.start()
    try:
        deadline = time.time() + 15
        while time.time() < deadline and not any(e.get("type") == "suspended" for e in events):
            time.sleep(0.05)
    finally:
        worker.stop()
        worker.join(15)
    failures = [e for e in events if e.get("type") == "failed"]
    assert failures, "no failed event"
    assert failures[0]["state"] == "failed_soft"


# -- 4. starting before the window opens should be feasible ---------------
def test_preflight_before_window_uses_the_coming_window(campaign):
    contacts = [{"email": f"early{i}@example.com"} for i in range(50)]
    cfg = campaign(contacts, quota=50, window_enabled=True, fixed_seconds=120)
    cfg.campaign.send_window.start = "09:00"
    cfg.campaign.send_window.end = "18:00"
    cfg.campaign.send_window.days = ["Mon", "Tue", "Wed", "Thu", "Fri"]
    now = datetime(2026, 10, 19, 8, 0, tzinfo=ZoneInfo("Europe/London"))  # Monday 08:00
    conn = db.connect(cfg.db_path)
    db.init_db(conn)
    report = preflight.run_preflight(cfg, conn, quota=50, now=now,
                                     today=date(2026, 10, 19))
    conn.close()
    assert report.feasibility is not None
    assert report.feasibility.remaining_seconds >= 9 * 3600
    assert report.feasibility.ok is True
    assert report.ok is True


def test_available_seconds_uses_the_next_window_when_closed():
    from dispatch.config import SendWindow
    window = SendWindow(start="09:00", end="18:00", timezone="Europe/London")
    tz = ZoneInfo("Europe/London")
    # before opening: judged against the nine hours that are about to open
    assert feasibility.available_seconds(datetime(2026, 10, 19, 8, 0, tzinfo=tz), window) == 9 * 3600
    # open: the rest of today
    assert feasibility.available_seconds(datetime(2026, 10, 19, 12, 0, tzinfo=tz), window) == 6 * 3600
    # closed: tomorrow's window, not zero
    assert feasibility.available_seconds(datetime(2026, 10, 19, 20, 0, tzinfo=tz), window) == 9 * 3600


# -- 5. write-back must honour date_order when comparing existing cells ---
def test_write_back_compares_dates_in_configured_order(campaign):
    cfg = campaign([{"email": "mdy@example.com", "last_advertised": "03/04/2026"}])
    cfg.campaign.date_order = "MDY"  # 03/04/2026 == 4 March 2026
    result = sheet_writer.write_back(
        cfg.sheet_file,
        [sheet_writer.Update(sheet_row=2, last_advertised=date(2026, 3, 15))],
        tab="Leads",
        date_order="MDY",
    )
    assert result.changes, "a later date was wrongly considered backwards"
    from openpyxl import load_workbook
    wb = load_workbook(cfg.sheet_file)
    value = wb["Leads"].cell(row=2, column=5).value
    wb.close()
    assert str(value)[:10] == "2026-03-15"


# -- 6. non-https markdown links are rejected, not silently dropped ------
def test_blocks_validator_rejects_http_markdown_links(campaign):
    doc = {
        "subject": "Hello",
        "blocks": [{"type": "paragraph", "text": "See [this](http://insecure.example)"}],
        "max_width_px": 600,
    }
    errors, _ = blocks.validate_document(doc, ["Email"])
    assert any("https" in e.lower() for e in errors), errors


# -- 7. starttls + implicit TLS at once is a config error ----------------
def test_smtp_cannot_use_starttls_and_ssl_together():
    payload = {
        "sender": {"smtp": {"host": "x", "port": 587, "starttls": True, "ssl": True}},
        "template": {"footer": {"postal_address": "x"}},
    }
    with pytest.raises(Exception):
        FullConfig.model_validate(payload)


# -- 8. re-adding a contact must not orphan its history -------------------
def test_contact_id_survives_removal_and_reimport(campaign):
    cfg = campaign([{"email": "a@example.com"}, {"email": "b@example.com"}])
    conn = db.connect(cfg.db_path)
    db.init_db(conn)
    from dispatch.sheet.reader import read_sheet
    first = read_sheet(cfg.sheet_file, "Leads")
    db.import_contacts(conn, first.contacts)
    b_id = db.contact_by_email(conn, "b@example.com")["id"]

    # b disappears from the sheet, then comes back
    from tests.conftest import write_workbook
    write_workbook(cfg.sheet_file, [{"email": "a@example.com"}])
    db.import_contacts(conn, read_sheet(cfg.sheet_file, "Leads").contacts)
    write_workbook(cfg.sheet_file, [{"email": "a@example.com"}, {"email": "b@example.com"}])
    db.import_contacts(conn, read_sheet(cfg.sheet_file, "Leads").contacts)

    assert db.contact_by_email(conn, "b@example.com")["id"] == b_id
    conn.close()


# -- 9. IMAP UIDVALIDITY resets the watermark ----------------------------
def test_uidvalidity_change_resets_the_watermark():
    assert monitor.uid_start(500, "111", "111") == 500
    assert monitor.uid_start(500, None, "111") == 500
    assert monitor.uid_start(500, "111", "222") == 0
