"""Shared fixtures: temporary campaigns, workbooks and a localhost SMTP sink.

No test touches the network: the SMTP server is aiosmtpd on 127.0.0.1 and the IMAP
side is exercised by parsing saved .eml files.
"""
from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

# Qt must pick the offscreen platform before any PySide6 import (UI smoke tests).
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from aiosmtpd.controller import Controller
from aiosmtpd.smtp import AuthResult
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from dispatch.config import FullConfig  # noqa: E402
from dispatch.mail import blocks  # noqa: E402

STARTERS = ROOT / "dispatch" / "templates" / "starters"
HEADERS = ["FirstName", "LastName", "Email", "Company", "LastAdvertised",
           "Replied", "Ignore", "LastContacted", "Status", "LastError", "Attempts",
           "Source"]


def write_workbook(path: Path, contacts: list[dict], headers: list[str] | None = None,
                   fill_row: bool = True) -> Path:
    headers = headers or HEADERS
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="1F3864")
    for contact in contacts:
        values = []
        for header in headers:
            key = header.lower()
            if key == "email":
                values.append(contact.get("email"))
            elif key == "firstname":
                values.append(contact.get("first_name"))
            elif key == "lastname":
                values.append(contact.get("last_name"))
            elif key == "company":
                values.append(contact.get("company"))
            elif key == "lastadvertised":
                values.append(contact.get("last_advertised"))
            elif key == "replied":
                values.append(contact.get("replied"))
            elif key == "ignore":
                values.append(contact.get("ignore"))
            elif key == "lastcontacted":
                values.append(contact.get("last_contacted"))
            elif key == "status":
                values.append(contact.get("status"))
            elif key == "lasterror":
                values.append(contact.get("last_error"))
            elif key == "attempts":
                values.append(contact.get("attempts"))
            else:
                values.append(contact.get(header, contact.get(key)))
        ws.append(values)
    if fill_row and len(contacts):
        ws.cell(row=2, column=1).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.column_dimensions["C"].width = 30
    wb.save(path)
    wb.close()
    return path


def make_config(tmp_path: Path, contacts: list[dict], *, quota: int = 200,
                dry_run: bool = False, window_enabled: bool = False,
                fixed_seconds: int = 0, smtp_port: int = 2525,
                shuffle: bool = False, headers: list[str] | None = None,
                daily_cap: int = 1000, hourly_cap: int = 1000,
                warmup: bool = False, frequency_days: int = 30,
                fail_rate: float = 1.0) -> FullConfig:
    base = tmp_path
    sheet = write_workbook(base / "leads.xlsx", contacts, headers=headers)
    tdir = base / "templates"
    tdir.mkdir(parents=True, exist_ok=True)
    doc = blocks.load_document(STARTERS / "personal_note.json")
    blocks.save_document(doc, tdir / "personal_note.json")
    (tdir / "personal_note.html").write_text(blocks.render_document_html(doc), encoding="utf-8")
    (tdir / "personal_note.txt").write_text(blocks.render_document_text(doc), encoding="utf-8")

    data = {
        "campaign": {
            "name": "Test campaign",
            "sheet_path": str(sheet),
            "sheet_tab": "Leads",
            "date_order": "DMY",
            "quota": quota,
            "frequency_days": frequency_days,
            "shuffle_order": shuffle,
            "send_window": {
                "enabled": window_enabled,
                "start": "00:00",
                "end": "23:59",
                "days": ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"],
                "timezone": "Europe/London",
            },
        },
        "throttle": {
            "mode": "fixed",
            "fixed_seconds": fixed_seconds,
            "random_seconds": [0, 0],
            "backoff_seconds": [1, 1, 1, 1],
        },
        "limits": {"daily_cap": daily_cap, "hourly_cap": hourly_cap},
        "safety": {"dry_run": dry_run,
                   "require_typed_confirmation_over": 50,
                   "stop_on_failure_rate": fail_rate},
        "warmup": {"enabled": warmup, "ramp": [30, 60, 100, 150, 200],
                   "first_send_date": None},
        "sender": {
            "id": "sam",
            "from_name": "Sam at Samlax",
            "from_email": "sam@samlax.net",
            "reply_to": "sam@samlax.net",
            "smtp": {"host": "127.0.0.1", "port": smtp_port, "starttls": False,
                     "timeout": 10, "ssl": False},
            "imap": {"host": "127.0.0.1", "port": 993, "ssl": True, "mailbox": "INBOX"},
            "secret_ref": "env:DISPATCH_TEST_PASSWORD",
            "max_per_connection": 50,
        },
        "template": {
            "kind": "blocks",
            "subject": "Quick question, {FirstName|default:\"there\"}",
            "html_path": str(tdir / "personal_note.html"),
            "text_path": str(tdir / "personal_note.txt"),
            "blocks_path": str(tdir / "personal_note.json"),
            "image_url": "",
            "footer": {
                "postal_address": "Samlax, 1 Test Street, London, E1 6AN",
                "unsubscribe_url": "https://samlax.net/u/{token}",
                "unsubscribe_mailto": "unsubscribe@samlax.net",
            },
        },
    }
    cfg = FullConfig.model_validate(data)
    cfg._base_dir = base
    return cfg


@pytest.fixture
def campaign(tmp_path):
    os.environ.setdefault("DISPATCH_TEST_PASSWORD", "test-password")
    return lambda contacts, **kw: make_config(tmp_path, contacts, **kw)


@pytest.fixture
def contacts_12():
    return [
        {"email": "alice@example.com", "first_name": "Alice", "last_name": "Perrin",
         "company": "Northwind", "last_advertised": date(2026, 1, 5)},
        {"email": "bob@example.com", "first_name": "Bob", "last_name": "Chen",
         "company": "Contoso"},
        {"email": "carol@example.com", "first_name": "Carol", "last_name": "Diaz",
         "company": "Fabrikam", "last_advertised": date(2026, 10, 12)},
        {"email": "dave@example.com", "first_name": "Dave", "last_name": "Evans",
         "ignore": "Y", "last_advertised": date(2025, 11, 2)},
        {"email": "erin@example.com", "first_name": "Erin", "last_name": "Fox",
         "replied": "Y", "last_advertised": date(2025, 12, 1)},
        {"email": "frank@example.com", "first_name": "Frank", "last_name": "Gray",
         "last_advertised": "03/04/2026"},
        {"email": "grace.hall@example", "first_name": "Grace", "last_name": "Hall"},
        {"email": "ALICE@example.com ", "first_name": "Alice", "last_name": "Perrin"},
        {"email": None, "company": "Umbrella"},
        {"email": "", "first_name": "Ivan", "company": "Soylent"},
        {"email": "judy@example.com", "first_name": "Judy", "last_name": "Klein",
         "last_advertised": date(2026, 9, 1)},
        {"email": "karl@example.com", "first_name": "Karl", "last_name": "Lang",
         "last_advertised": date(2026, 10, 1)},
    ]


# --------------------------------------------------------------------------
# localhost SMTP
# --------------------------------------------------------------------------
class Sink:
    def __init__(self):
        self.messages: list[dict] = []
        self.fail_next = 0
        self.fail_code = "421"
        self.fail_text = "4.7.28 Too many messages"
        self.recipients: list[str] = []


class _Handler:
    def __init__(self, sink: Sink):
        self.sink = sink

    async def handle_RCPT(self, server, session, envelope, address, rcpt_options):
        envelope.rcpt_tos.append(address)
        self.sink.recipients.append(address)
        return "250 OK"

    async def handle_DATA(self, server, session, envelope):
        if self.sink.fail_next:
            self.sink.fail_next -= 1
            return f"{self.sink.fail_code} {self.sink.fail_text}"
        self.sink.messages.append({
            "from": envelope.mail_from,
            "to": list(envelope.rcpt_tos),
            "data": envelope.content,
        })
        return "250 Message accepted for delivery"


def _authenticator(server, session, envelope, mechanism, auth_data):
    return AuthResult(success=True)


def _free_port() -> int:
    """Controller(port=0) cannot start on Windows (its readiness probe dials port 0)."""
    import socket
    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])
    finally:
        sock.close()


@pytest.fixture
def smtp_sink():
    sink = Sink()
    port = _free_port()
    controller = Controller(_Handler(sink), hostname="127.0.0.1", port=port,
                            authenticator=_authenticator, auth_require_tls=False)
    controller.start()
    try:
        sink.port = port
        yield sink
    finally:
        controller.stop()


@pytest.fixture
def fixed_today():
    return date(2026, 10, 19)
