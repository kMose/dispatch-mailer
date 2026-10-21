"""Regenerate sample_leads.xlsx and the promo template artifacts.

Run from the project root:  python tools/make_fixtures.py
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from dispatch.mail import blocks  # noqa: E402

HEADERS = ["FirstName", "LastName", "Email", "Company", "LastAdvertised",
           "Replied", "Ignore", "LastContacted", "Status", "LastError", "Attempts",
           "Source", "Notes"]

ROWS = [
    ["Alice", "Perrin", "alice@example.com", "Northwind", date(2026, 1, 5), None, None,
     date(2026, 1, 5), None, None, 1, "referral", "warm lead"],
    ["Bob", "Chen", "bob@example.com", "Contoso", None, None, None, None, None, None,
     None, "webinar", "never contacted"],
    ["Carol", "Diaz", "carol@example.com", "Fabrikam", date(2026, 10, 12), None, None,
     date(2026, 10, 12), "sent", None, 1, "webinar", ""],
    ["Dave", "Evans", "dave@example.com", "Tailspin", date(2025, 11, 2), None, "Y",
     date(2025, 11, 2), "sent", None, 2, "cold list", "asked us to pause"],
    ["Erin", "Fox", "erin@example.com", "Adventure Works", date(2025, 12, 1), "Y", None,
     date(2025, 12, 1), "sent", None, 1, "referral", "replied in January"],
    ["Frank", "Gray", "frank@example.com", "Globex", "03/04/2026", None, None,
     "03/04/2026", None, None, None, "conference", "ambiguous date on purpose"],
    ["Grace", "Hall", "grace.hall@example", "Initech", date(2025, 9, 9), None, None,
     date(2025, 9, 9), None, None, 1, "cold list", "domain has no dot"],
    ["Alice", "Perrin", "ALICE@example.com ", "Northwind", date(2026, 1, 5), None, None,
     date(2026, 1, 5), None, None, 1, "referral", "duplicate of row 1"],
    [None, None, None, "Umbrella", None, None, None, None, None, None, None,
     "manual", "no email at all"],
    ["Ivan", "Jones", None, "Soylent", None, None, None, None, None, None, None,
     "manual", "blank email cell"],
    ["Judy", "Klein", "judy@example.com", "Hooli", date(2026, 9, 1), None, None,
     date(2026, 9, 1), "sent", None, 1, "referral", "45+ day gap"],
    ["Karl", "Lang", "karl@example.com", "Vandelay", date(2026, 10, 1), None, None,
     date(2026, 10, 1), "sent", None, 1, "webinar", "recently contacted"],
]


def build_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="1F3864")
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = max(14, len(HEADERS[col - 1]) + 4)

    for r in range(2, len(ROWS) + 2):
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=8).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.cell(row=2, column=1).fill = PatternFill("solid", fgColor="FFF2CC")

    wb.save(path)
    wb.close()


def build_promo_templates(templates: Path) -> None:
    doc = blocks.load_document(ROOT / "dispatch" / "templates" / "starters" / "promo_card.json")
    doc["name"] = "October promo"
    blocks.save_document(doc, templates / "promo.json")
    (templates / "promo.html").write_text(blocks.render_document_html(doc), encoding="utf-8")
    (templates / "promo.txt").write_text(blocks.render_document_text(doc), encoding="utf-8")


if __name__ == "__main__":
    build_workbook(ROOT / "sample_leads.xlsx")
    build_promo_templates(ROOT / "templates")
    print("wrote sample_leads.xlsx and templates/promo.{json,html,txt}")
